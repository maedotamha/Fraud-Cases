"""
Build Interim 2 report — reads real model results from CSV and JSON,
then builds the Word .docx and Medium HTML with actual numbers.
Run after modeling.ipynb has fully executed.
"""
import pandas as pd
import numpy as np
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Paths ────────────────────────────────────────────────────────────────────
BASE   = Path(".")
VIS    = BASE / "report" / "visuals"
MODELS = BASE / "models"
PROC   = BASE / "data" / "processed"
OUT    = BASE / "report"

# ── Load results ─────────────────────────────────────────────────────────────
df = pd.read_csv(PROC / "model_comparison.csv")
with open(MODELS / "best_params.json") as f:
    best_params = json.load(f)

ecomm  = df[df["Dataset"] == "E-Commerce"].reset_index(drop=True)
cc     = df[df["Dataset"] == "Credit Card"].reset_index(drop=True)
lr_ec  = ecomm[ecomm["Model"].str.contains("Logistic")].iloc[0]
xgb_ec = ecomm[ecomm["Model"].str.contains("XGBoost")].iloc[0]
lr_cc  = cc[cc["Model"].str.contains("Logistic")].iloc[0]
xgb_cc = cc[cc["Model"].str.contains("XGBoost")].iloc[0]

print("Results loaded:")
print(df[["Dataset","Model","AUCPR","F1","TP","FP","FN","TN"]].to_string(index=False))

# ── Colour constants ─────────────────────────────────────────────────────────
GREEN  = RGBColor(0x1A, 0x89, 0x17)
RED    = RGBColor(0xC0, 0x39, 0x2B)
BLUE   = RGBColor(0x0D, 0x6E, 0xFD)
GREY   = RGBColor(0x9B, 0x9B, 0x9B)
DARK   = RGBColor(0x1A, 0x1A, 0x1A)
ORANGE = RGBColor(0xB0, 0x7D, 0x1A)

# ── Docx helpers ─────────────────────────────────────────────────────────────
def shade_para(para, hex_color="FFF8E7"):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color); pPr.append(shd)

def shade_cell(cell, hex_color="F7F7F7"):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color); tcPr.append(shd)

def add_table(doc, headers, rows, col_widths=None, highlight=None, hdr_color="2D6A4F"):
    highlight = highlight or []
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = "Table Grid"
    hc = t.rows[0].cells
    for i, h in enumerate(headers):
        hc[i].text = h
        p = hc[i].paragraphs[0]
        p.runs[0].bold = True; p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(hc[i], hdr_color)
    for ri, row in enumerate(rows):
        rc = t.rows[ri+1].cells
        hi = ri in highlight
        for ci, val in enumerate(row):
            rc[ci].text = str(val)
            p = rc[ci].paragraphs[0]
            p.runs[0].font.size = Pt(9.5)
            if hi:
                shade_cell(rc[ci], "FFF3CD")
                p.runs[0].bold = True
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in t.rows: row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t

def add_image(doc, fname, width=6.0, caption=None):
    img_path = VIS / fname
    if not img_path.exists():
        doc.add_paragraph(f"[Figure: {fname} — not yet generated]")
        return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(img_path), width=Inches(width))
    if caption:
        c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].font.size = Pt(9); c.runs[0].italic = True
        c.runs[0].font.color.rgb = GREY
    doc.add_paragraph()

def callout(doc, label, text, bg="FFF8E7", label_color=None):
    p = doc.add_paragraph(); shade_para(p, bg)
    r = p.add_run(label + "  "); r.bold = True
    if label_color: r.font.color.rgb = label_color
    p.add_run(text); doc.add_paragraph()

def h1(doc, text):
    h = doc.add_heading(text, level=1); h.runs[0].font.color.rgb = GREEN; return h
def h2(doc, text):
    h = doc.add_heading(text, level=2); h.runs[0].font.color.rgb = DARK;  return h
def h3(doc, text):
    h = doc.add_heading(text, level=3); return h
def body(doc, text):
    p = doc.add_paragraph(text)
    for run in p.runs: run.font.size = Pt(10.5)
    return p

# ════════════════════════════════════════════════════════════════════════════
# BUILD WORD DOCUMENT
# ════════════════════════════════════════════════════════════════════════════
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.8)

# ── Cover ────────────────────────────────────────────────────────────────────
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("From Data to Decisions: Training Fraud Detection Models")
r.bold = True; r.font.size = Pt(24); r.font.color.rgb = GREEN

p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Adey Innovations Inc. — Interim Report 2"); r2.font.size = Pt(15); r2.italic = True; r2.font.color.rgb = GREY

p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = p3.add_run("Task 2 — Model Building, Training & Evaluation"); r3.font.size = Pt(12); r3.bold = True

p4 = doc.add_paragraph(); p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
r4 = p4.add_run("Maedot Amha  ·  June 2026"); r4.font.size = Pt(10); r4.font.color.rgb = GREY
doc.add_paragraph()

lp = doc.add_paragraph(); shade_para(lp, "F2FDF2")
lp.add_run(
    "This report presents two complete fraud detection pipelines — one for e-commerce transactions, "
    "one for bank credit cards. Each pipeline trains a Logistic Regression baseline and a tuned "
    "XGBoost model, evaluated on AUC-PR, F1-Score, and Confusion Matrix via Stratified K-Fold "
    "cross-validation. All models are saved and ready for SHAP explainability in Task 3."
).font.size = Pt(11)
doc.add_page_break()

# ── Section 1: Overview ──────────────────────────────────────────────────────
h1(doc, "1.  Pipeline Overview")
body(doc,
    "The two datasets are treated as independent modeling problems. They differ in feature structure, "
    "fraud rate, and dominant signal type — a shared pipeline would mask these differences.")
add_table(doc,
    headers=["", "E-Commerce Pipeline", "Credit Card Pipeline"],
    rows=[
        ["Training rows (post-SMOTE)", "219,136 (balanced 1:1)", "453,204 (balanced 1:1)"],
        ["Test rows",                   "30,255 (original ratio)", "56,746 (original ratio)"],
        ["Features",                    "194 (OHE + engineered)", "30 (V1–V28, Amount, Time)"],
        ["Fraud rate (test)",            "9.36%",                  "0.17%"],
        ["Target column",               "'class'",                 "'Class'"],
        ["Baseline model",              "Logistic Regression",     "Logistic Regression"],
        ["Ensemble model",              "XGBoost (tuned)",         "XGBoost (tuned)"],
        ["Primary metric",              "AUC-PR, F1",              "AUC-PR, F1"],
    ],
    col_widths=[1.9, 2.3, 2.3]
)

# ── Section 2: Baseline ──────────────────────────────────────────────────────
h1(doc, "2.  Baseline: Logistic Regression")
body(doc,
    "Logistic Regression with class_weight='balanced' and the saga solver (handles large sparse "
    "feature matrices efficiently). This establishes the minimum performance bar any ensemble "
    "model must beat to justify its complexity.")
add_table(doc,
    headers=["Dataset", "AUC-PR", "F1-Score", "TP", "FP", "FN", "TN"],
    rows=[
        ["E-Commerce",  lr_ec["AUCPR"],  lr_ec["F1"],  int(lr_ec["TP"]),  int(lr_ec["FP"]),  int(lr_ec["FN"]),  int(lr_ec["TN"])],
        ["Credit Card", lr_cc["AUCPR"],  lr_cc["F1"],  int(lr_cc["TP"]),  int(lr_cc["FP"]),  int(lr_cc["FN"]),  int(lr_cc["TN"])],
    ],
    col_widths=[1.3, 0.9, 0.9, 0.8, 0.8, 0.8, 1.2]
)

add_image(doc, "figM3_confusion_matrices.png", width=6.2,
    caption="Figure 1 — Confusion matrices for all models. Row = true label, Column = predicted label. Percentages are row-normalised.")

# ── Section 3: XGBoost ───────────────────────────────────────────────────────
h1(doc, "3.  Ensemble Model: XGBoost with Hyperparameter Tuning")
h2(doc, "3.1  Why XGBoost?")
body(doc,
    "XGBoost builds an ensemble of decision trees using gradient boosting. Each tree corrects "
    "the errors of the previous one, making it highly effective at capturing non-linear fraud "
    "patterns that Logistic Regression cannot model with linear boundaries.")
add_table(doc,
    headers=["Property", "Logistic Regression", "XGBoost"],
    rows=[
        ["Decision boundary", "Linear only",          "Non-linear (ensemble trees)"],
        ["Handles interactions", "No",                "Yes — automatically captures feature interactions"],
        ["Imbalance handling", "class_weight param",  "scale_pos_weight + SMOTE"],
        ["Interpretability",   "Direct coefficients", "SHAP values (Task 3)"],
        ["Training speed",     "Fast",                "Slower, but parallelizable"],
    ],
    col_widths=[1.5, 2.0, 2.8]
)

h2(doc, "3.2  Hyperparameter Tuning — RandomizedSearchCV")
body(doc,
    "RandomizedSearchCV with 12 random configurations across the search space below, "
    "evaluated via 5-fold stratified CV scored on average_precision (AUC-PR):")
ec_params  = best_params.get("ecommerce", {})
cc_params  = best_params.get("creditcard", {})
param_rows = []
all_keys = sorted(set(list(ec_params.keys()) + list(cc_params.keys())))
for k in all_keys:
    param_rows.append([k, str(ec_params.get(k, "—")), str(cc_params.get(k, "—"))])
add_table(doc,
    headers=["Parameter", "Best Value — E-Commerce", "Best Value — Credit Card"],
    rows=param_rows,
    col_widths=[1.8, 2.1, 2.1]
)

h2(doc, "3.3  XGBoost Test Set Results")
add_table(doc,
    headers=["Dataset", "AUC-PR", "F1-Score", "TP", "FP", "FN", "TN"],
    rows=[
        ["E-Commerce",  xgb_ec["AUCPR"], xgb_ec["F1"], int(xgb_ec["TP"]), int(xgb_ec["FP"]), int(xgb_ec["FN"]), int(xgb_ec["TN"])],
        ["Credit Card", xgb_cc["AUCPR"], xgb_cc["F1"], int(xgb_cc["TP"]), int(xgb_cc["FP"]), int(xgb_cc["FN"]), int(xgb_cc["TN"])],
    ],
    col_widths=[1.3, 0.9, 0.9, 0.8, 0.8, 0.8, 1.2],
    highlight=[0, 1]
)

# ── Section 4: CV ────────────────────────────────────────────────────────────
h1(doc, "4.  Stratified K-Fold Cross-Validation (k=5)")
body(doc,
    "Standard K-Fold risks concentrating all fraud cases in one fold on severely imbalanced data. "
    "Stratified K-Fold guarantees each fold preserves the class ratio. "
    "SMOTE was applied only to training folds — never to the validation fold.")
add_table(doc,
    headers=["Dataset", "CV AUC-PR (mean ± std)", "CV F1 (mean ± std)", "Interpretation"],
    rows=[
        ["E-Commerce",  xgb_ec["CV_AUCPR"], xgb_ec["CV_F1"],  "Stable across folds if std < 0.02"],
        ["Credit Card", xgb_cc["CV_AUCPR"], xgb_cc["CV_F1"],  "Stable across folds if std < 0.02"],
    ],
    col_widths=[1.3, 2.0, 1.8, 2.1]
)
add_image(doc, "figM4_cv_results.png", width=6.2,
    caption="Figure 2 — Per-fold AUC-PR and F1-Score for XGBoost (k=5). Dashed lines = mean. Narrow spread = reliable generalisation.")

# ── Section 5: PR curves ─────────────────────────────────────────────────────
h1(doc, "5.  Precision-Recall Curves")
body(doc,
    "The Precision-Recall curve plots the trade-off between catching more fraud (high recall) "
    "and avoiding false alarms (high precision). The area under this curve (AUC-PR) is our primary "
    "evaluation metric — it is far more informative than ROC-AUC on imbalanced data.")
add_image(doc, "figM2_pr_curves.png", width=6.2,
    caption="Figure 3 — PR curves for both models on both datasets. XGBoost (dashed) dominates Logistic Regression. The grey dotted line is the random baseline.")

callout(doc,
    "📈  Reading the PR curve:",
    "A model hugging the top-right corner is ideal — high precision and high recall simultaneously. "
    "The steeper the initial rise, the better the model ranks true fraud cases at the top of its "
    "score distribution. XGBoost achieves this on both datasets.",
    bg="FFF8E7", label_color=ORANGE)

# ── Section 6: Model comparison ──────────────────────────────────────────────
h1(doc, "6.  Full Model Comparison")
add_table(doc,
    headers=["Dataset", "Model", "AUC-PR", "F1", "CV AUC-PR", "CV F1", "FN (missed fraud)"],
    rows=[
        ["E-Commerce",  "Logistic Regression", lr_ec["AUCPR"],  lr_ec["F1"],  "—",               "—",               int(lr_ec["FN"])],
        ["E-Commerce",  "XGBoost (tuned) ★",   xgb_ec["AUCPR"], xgb_ec["F1"], xgb_ec["CV_AUCPR"], xgb_ec["CV_F1"],   int(xgb_ec["FN"])],
        ["Credit Card", "Logistic Regression", lr_cc["AUCPR"],  lr_cc["F1"],  "—",               "—",               int(lr_cc["FN"])],
        ["Credit Card", "XGBoost (tuned) ★",   xgb_cc["AUCPR"], xgb_cc["F1"], xgb_cc["CV_AUCPR"], xgb_cc["CV_F1"],   int(xgb_cc["FN"])],
    ],
    col_widths=[1.2, 1.6, 0.8, 0.7, 1.7, 1.5, 1.3],
    highlight=[1, 3]
)

# ── Section 7: Feature importance ────────────────────────────────────────────
h1(doc, "7.  Feature Importance — XGBoost Built-In")
body(doc,
    "XGBoost's gain-based feature importance measures how much each feature contributes to reducing "
    "prediction error across all trees. This provides a model-internal view before the more "
    "rigorous SHAP analysis in Task 3.")
add_image(doc, "figM5_feature_importance.png", width=6.2,
    caption="Figure 4 — Top 10 features by XGBoost gain importance. Red bar = most important feature. "
            "E-commerce: time_since_signup dominates. Credit card: V17/V14 dominate as expected from EDA.")

callout(doc,
    "🔍  Preview of SHAP findings:",
    "Built-in feature importance confirms our EDA signal hierarchy: time_since_signup leads in "
    "e-commerce; V17 and V14 lead in credit card. Task 3 will validate whether these are "
    "consistent across individual predictions or context-dependent using SHAP force plots.",
    bg="F7F9FF", label_color=BLUE)

# ── Section 8: Model selection ───────────────────────────────────────────────
h1(doc, "8.  Model Selection — Justification")
reasons = [
    ("Performance", f"XGBoost achieves higher AUC-PR ({xgb_ec['AUCPR']} vs {lr_ec['AUCPR']} on e-commerce; "
     f"{xgb_cc['AUCPR']} vs {lr_cc['AUCPR']} on credit card) and higher F1 on both datasets."),
    ("Fewer Missed Frauds", f"XGBoost produces fewer false negatives ({int(xgb_ec['FN'])} vs {int(lr_ec['FN'])} "
     f"on e-commerce), meaning fewer fraudulent transactions slip through undetected."),
    ("Cross-Validation Stability", "Low standard deviation in CV AUC-PR confirms the model generalises "
     "reliably across unseen data folds — not overfitting to the SMOTE-augmented training set."),
    ("Non-Linear Fraud Patterns", "Fraud behaviour involves complex feature interactions "
     "(e.g. high purchase_value AND new account AND high-risk country). "
     "XGBoost captures these; Logistic Regression cannot."),
    ("SHAP Compatibility", "Despite being a black-box ensemble, XGBoost integrates directly with "
     "SHAP's TreeExplainer — providing fast, exact feature attribution for regulatory explainability."),
]
for i, (title, text) in enumerate(reasons, 1):
    p = doc.add_paragraph(); shade_para(p, "F7F9FF")
    r = p.add_run(f"Reason {i}: {title}\n"); r.bold = True; r.font.color.rgb = BLUE
    p.add_run(text).font.size = Pt(10.5)
    doc.add_paragraph()

# ── Section 9: Next Steps ─────────────────────────────────────────────────────
h1(doc, "9.  Next Steps — Task 3 (Final Submission)")
add_table(doc,
    headers=["Task", "Detail", "Deliverable"],
    rows=[
        ["SHAP summary plot",        "Global feature importance via SHAP values on test set", "Bar + beeswarm plot per dataset"],
        ["SHAP force plots",         "3 individual predictions: TP, FP, FN",                 "Force plot per case with annotation"],
        ["Built-in vs SHAP compare", "Validate gain importance against SHAP rank",            "Comparison table"],
        ["Business recommendations", "≥3 actionable insights connected to SHAP values",       "Recommendation list with SHAP evidence"],
        ["Final report",             "End-to-end blog/PDF covering all three tasks",           "Published report + final GitHub push"],
    ],
    col_widths=[1.6, 2.6, 2.2]
)

# ── Footer ────────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
r = p.add_run("GitHub: https://github.com/maedotamha/Fraud-Cases   ·   "
              "Adey Innovations Inc.  ·  Interim Report 2  ·  June 2026")
r.font.size = Pt(8.5); r.font.color.rgb = GREY

out_path = OUT / "interim2_report.docx"
doc.save(str(out_path))
print(f"\nWord report saved: {out_path}")

# ════════════════════════════════════════════════════════════════════════════
# BUILD MEDIUM HTML
# ════════════════════════════════════════════════════════════════════════════
# Helper for inline base64 images so the HTML is self-contained
import base64

def img_tag(fname, alt, caption):
    path = VIS / fname
    if not path.exists():
        return f'<p style="text-align:center;color:#aaa;font-style:italic">[{fname} — awaiting training]</p>'
    data = base64.b64encode(path.read_bytes()).decode()
    ext  = fname.split(".")[-1]
    return (f'<div class="fig"><img src="data:image/{ext};base64,{data}" alt="{alt}">'
            f'<div class="fig-caption">{caption}</div></div>')

html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>From Data to Decisions — Interim Report 2</title>
<style>
  body{{font-family:'Georgia',serif;max-width:740px;margin:60px auto;color:#292929;line-height:1.85;font-size:18px}}
  h1{{font-size:40px;font-weight:700;line-height:1.2;margin-bottom:8px}}
  h2{{font-size:26px;font-weight:700;margin-top:52px;margin-bottom:14px;border-bottom:2px solid #e8e8e8;padding-bottom:8px}}
  h3{{font-size:19px;font-weight:700;margin-top:32px;margin-bottom:6px}}
  .subtitle{{font-size:20px;color:#6b6b6b;margin-bottom:6px;font-style:italic}}
  .meta{{font-size:14px;color:#9b9b9b;margin-bottom:44px}}
  .lead{{font-size:20px;line-height:1.65;color:#3d3d3d;margin-bottom:36px;border-left:4px solid #1a8917;padding-left:20px}}
  .callout{{border-left:4px solid #f5a623;background:#fff8e7;padding:18px 22px;margin:28px 0;border-radius:4px}}
  .callout strong{{display:block;font-size:13px;text-transform:uppercase;letter-spacing:.06em;color:#b07d1a;margin-bottom:6px}}
  .callout-red{{border-left-color:#e74c3c;background:#fff2f2}}.callout-red strong{{color:#c0392b}}
  .callout-blue{{border-left-color:#0d6efd;background:#f0f4ff}}.callout-blue strong{{color:#0d6efd}}
  table{{width:100%;border-collapse:collapse;margin:22px 0;font-size:15px}}
  th{{background:#2d6a4f;color:#fff;text-align:left;padding:9px 13px;font-weight:700}}
  td{{padding:9px 13px;border-bottom:1px solid #eee}}
  tr:last-child td{{border-bottom:none}}
  .hl{{background:#fff3cd;font-weight:700}}
  code{{background:#f4f4f4;padding:2px 6px;border-radius:3px;font-family:'Courier New',monospace;font-size:14px}}
  pre{{background:#1e1e1e;color:#d4d4d4;padding:22px;border-radius:8px;overflow-x:auto;font-size:13px;line-height:1.65;margin:24px 0}}
  pre .kw{{color:#569cd6}} pre .fn{{color:#dcdcaa}} pre .st{{color:#ce9178}} pre .cm{{color:#6a9955}}
  blockquote{{border-left:3px solid #292929;padding-left:18px;margin:28px 0;font-style:italic;color:#555;font-size:19px}}
  .finding{{background:#f7f9ff;border:1px solid #cfe2ff;border-radius:8px;padding:18px 22px;margin:14px 0}}
  .finding-n{{font-size:12px;text-transform:uppercase;font-weight:700;color:#0d6efd;letter-spacing:.08em}}
  .tag{{display:inline-block;background:#e8f4f8;color:#1a6b8a;font-size:13px;padding:4px 10px;border-radius:20px;margin-right:6px}}
  .stat-row{{display:flex;gap:12px;flex-wrap:wrap;margin:20px 0}}
  .stat{{flex:1;min-width:140px;background:#f2fdf2;border:1px solid #c3e6cb;border-radius:8px;padding:14px 18px;text-align:center}}
  .stat .n{{font-size:32px;font-weight:700;color:#1a8917;display:block}} .stat .l{{font-size:12px;color:#555}}
  .stat.red{{background:#fff2f2;border-color:#f5c6cb}} .stat.red .n{{color:#c0392b}}
  .fig{{text-align:center;margin:28px 0}} .fig img{{max-width:100%;border-radius:6px;box-shadow:0 2px 12px rgba(0,0,0,.12)}}
  .fig-caption{{font-size:13px;color:#888;margin-top:8px;font-style:italic}}
  hr{{border:none;border-top:1px solid #e8e8e8;margin:44px 0}}
  .footer{{font-size:13px;color:#9b9b9b;margin-top:56px;padding-top:20px;border-top:1px solid #e8e8e8}}
</style>
</head>
<body>
<p><span class="tag">Fraud Detection</span><span class="tag">XGBoost</span><span class="tag">Machine Learning</span><span class="tag">FinTech</span></p>
<h1>From Data to Decisions</h1>
<p class="subtitle">Training & Evaluating Fraud Detection Models — Interim Report 2</p>
<p class="meta">Maedot Amha &nbsp;·&nbsp; Adey Innovations Inc. &nbsp;·&nbsp; June 2026 &nbsp;·&nbsp; 12 min read</p>
<p class="lead">With clean, feature-engineered data in hand, the question becomes: which model catches the most fraud while minimising false alarms? This report documents two complete modeling pipelines — e-commerce and bank credit cards — comparing Logistic Regression against a tuned XGBoost ensemble.</p>
<hr>

<h2>1. The Setup: Two Independent Pipelines</h2>
<p>Treating both datasets with a shared pipeline would be a mistake. Their feature structures, fraud rates, and dominant signals are fundamentally different:</p>
<table>
<thead><tr><th></th><th>E-Commerce</th><th>Credit Card</th></tr></thead>
<tbody>
<tr><td>Training rows (post-SMOTE)</td><td>219,136 (balanced 1:1)</td><td>453,204 (balanced 1:1)</td></tr>
<tr><td>Test rows</td><td>30,255 (original ratio)</td><td>56,746 (original ratio)</td></tr>
<tr><td>Features</td><td>194 (OHE + engineered)</td><td>30 (V1–V28, Amount, Time)</td></tr>
<tr><td>Fraud rate (test)</td><td>9.36%</td><td>0.17%</td></tr>
<tr><td>Primary signal</td><td>time_since_signup, country</td><td>V17, V14, V12</td></tr>
</tbody>
</table>
<div class="callout callout-red">
<strong>⚠️ Why accuracy is still wrong here</strong>
A model predicting all-legitimate on the credit card test set achieves 99.83% accuracy. We evaluate exclusively on <strong>AUC-PR and F1-Score</strong> — metrics that penalise missed fraud cases directly.
</div>
<hr>

<h2>2. Baseline: Logistic Regression</h2>
<p>Logistic Regression with <code>class_weight='balanced'</code> and the <code>saga</code> solver — efficient on large sparse matrices from one-hot encoding. This establishes the minimum performance bar.</p>
<table>
<thead><tr><th>Dataset</th><th>AUC-PR</th><th>F1-Score</th><th>TP</th><th>FP</th><th>FN</th><th>TN</th></tr></thead>
<tbody>
<tr><td>E-Commerce</td><td>{lr_ec["AUCPR"]}</td><td>{lr_ec["F1"]}</td><td>{int(lr_ec["TP"]):,}</td><td>{int(lr_ec["FP"]):,}</td><td>{int(lr_ec["FN"]):,}</td><td>{int(lr_ec["TN"]):,}</td></tr>
<tr><td>Credit Card</td><td>{lr_cc["AUCPR"]}</td><td>{lr_cc["F1"]}</td><td>{int(lr_cc["TP"]):,}</td><td>{int(lr_cc["FP"]):,}</td><td>{int(lr_cc["FN"]):,}</td><td>{int(lr_cc["TN"]):,}</td></tr>
</tbody>
</table>
<hr>

<h2>3. Ensemble Model: XGBoost with Hyperparameter Tuning</h2>
<h3>Why XGBoost over Logistic Regression?</h3>
<table>
<thead><tr><th>Property</th><th>Logistic Regression</th><th>XGBoost</th></tr></thead>
<tbody>
<tr><td>Decision boundary</td><td>Linear only</td><td>Non-linear — captures complex patterns</td></tr>
<tr><td>Feature interactions</td><td>Must be manually engineered</td><td>Discovered automatically</td></tr>
<tr><td>Imbalance handling</td><td>class_weight only</td><td>scale_pos_weight + SMOTE</td></tr>
<tr><td>Interpretability</td><td>Direct coefficients</td><td>SHAP (Task 3)</td></tr>
</tbody>
</table>

<h3>Hyperparameter Search (RandomizedSearchCV, 12 iterations, 5-fold CV)</h3>
<pre>
<span class="cm"># Search space</span>
param_dist = {{
    <span class="st">'n_estimators'</span>:     [200, 400],
    <span class="st">'max_depth'</span>:        [3, 5, 7],
    <span class="st">'learning_rate'</span>:    [0.05, 0.1, 0.2],
    <span class="st">'subsample'</span>:        [0.7, 0.9],
    <span class="st">'colsample_bytree'</span>: [0.7, 0.9],
    <span class="st">'min_child_weight'</span>: [1, 3],
}}
<span class="cm"># Scored on average_precision (= AUC-PR)</span>
search = <span class="fn">RandomizedSearchCV</span>(xgb, param_dist, n_iter=12, scoring=<span class="st">'average_precision'</span>, cv=skf)
</pre>

<p><strong>Best parameters found:</strong></p>
<table>
<thead><tr><th>Parameter</th><th>E-Commerce</th><th>Credit Card</th></tr></thead>
<tbody>
{"".join(f"<tr><td><code>{k}</code></td><td>{ec_params.get(k,'—')}</td><td>{cc_params.get(k,'—')}</td></tr>" for k in all_keys)}
</tbody>
</table>

<h3>XGBoost Test Results</h3>
<table>
<thead><tr><th>Dataset</th><th>AUC-PR</th><th>F1-Score</th><th>TP</th><th>FP</th><th>FN</th><th>TN</th></tr></thead>
<tbody>
<tr class="hl"><td>E-Commerce</td><td>{xgb_ec["AUCPR"]}</td><td>{xgb_ec["F1"]}</td><td>{int(xgb_ec["TP"]):,}</td><td>{int(xgb_ec["FP"]):,}</td><td>{int(xgb_ec["FN"]):,}</td><td>{int(xgb_ec["TN"]):,}</td></tr>
<tr class="hl"><td>Credit Card</td><td>{xgb_cc["AUCPR"]}</td><td>{xgb_cc["F1"]}</td><td>{int(xgb_cc["TP"]):,}</td><td>{int(xgb_cc["FP"]):,}</td><td>{int(xgb_cc["FN"]):,}</td><td>{int(xgb_cc["TN"]):,}</td></tr>
</tbody>
</table>
<hr>

<h2>4. Cross-Validation Results (Stratified K-Fold, k=5)</h2>
<p>Standard K-Fold risks concentrating all fraud cases in one fold. Stratified K-Fold guarantees each fold preserves the class ratio. <strong>SMOTE was applied only to training folds</strong> — the validation fold is always left at the original class distribution.</p>
<table>
<thead><tr><th>Dataset</th><th>CV AUC-PR (mean ± std)</th><th>CV F1 (mean ± std)</th></tr></thead>
<tbody>
<tr><td>E-Commerce XGBoost</td><td>{xgb_ec["CV_AUCPR"]}</td><td>{xgb_ec["CV_F1"]}</td></tr>
<tr><td>Credit Card XGBoost</td><td>{xgb_cc["CV_AUCPR"]}</td><td>{xgb_cc["CV_F1"]}</td></tr>
</tbody>
</table>
{img_tag("figM4_cv_results.png", "CV fold scores", "Figure 2 — Per-fold AUC-PR and F1 for XGBoost (k=5). Low std dev = stable generalisation.")}

<hr>
<h2>5. Precision-Recall Curves</h2>
<p>The PR curve plots precision vs recall at every decision threshold. AUC-PR summarises it in a single number — it directly penalises models that miss fraud cases or produce too many false alarms.</p>
{img_tag("figM2_pr_curves.png", "PR curves", "Figure 3 — PR curves for both models on both datasets. XGBoost (dashed) dominates throughout. The grey dotted line is the random classifier baseline.")}

<blockquote>A model hugging the top-right corner achieves high precision and high recall simultaneously — catching most fraud while generating few false alerts.</blockquote>
<hr>

<h2>6. Confusion Matrices</h2>
{img_tag("figM3_confusion_matrices.png", "Confusion matrices", "Figure 1 — Normalised confusion matrices (row = true label, column = predicted). Raw counts shown in grey.")}
<hr>

<h2>7. Feature Importance (Built-In)</h2>
{img_tag("figM5_feature_importance.png", "Feature importance", "Figure 4 — Top 10 XGBoost gain-importance features per dataset. Red = highest. E-commerce: time_since_signup leads. Credit card: V17/V14 lead.")}
<div class="callout callout-blue">
<strong>🔍 Preview of Task 3 SHAP Findings</strong>
Built-in importance confirms our EDA signal hierarchy. However, gain importance can overstate features with many splits. SHAP's TreeExplainer will provide more reliable, unbiased attribution — and explain individual predictions (TP, FP, FN) for regulatory justification.
</div>
<hr>

<h2>8. Full Model Comparison</h2>
<table>
<thead><tr><th>Dataset</th><th>Model</th><th>AUC-PR</th><th>F1</th><th>CV AUC-PR</th><th>FN (missed)</th></tr></thead>
<tbody>
<tr><td>E-Commerce</td><td>Logistic Regression</td><td>{lr_ec["AUCPR"]}</td><td>{lr_ec["F1"]}</td><td>—</td><td>{int(lr_ec["FN"]):,}</td></tr>
<tr class="hl"><td>E-Commerce</td><td>XGBoost (tuned) ★</td><td>{xgb_ec["AUCPR"]}</td><td>{xgb_ec["F1"]}</td><td>{xgb_ec["CV_AUCPR"]}</td><td>{int(xgb_ec["FN"]):,}</td></tr>
<tr><td>Credit Card</td><td>Logistic Regression</td><td>{lr_cc["AUCPR"]}</td><td>{lr_cc["F1"]}</td><td>—</td><td>{int(lr_cc["FN"]):,}</td></tr>
<tr class="hl"><td>Credit Card</td><td>XGBoost (tuned) ★</td><td>{xgb_cc["AUCPR"]}</td><td>{xgb_cc["F1"]}</td><td>{xgb_cc["CV_AUCPR"]}</td><td>{int(xgb_cc["FN"]):,}</td></tr>
</tbody>
</table>
<hr>

<h2>9. Model Selection: XGBoost for Both Datasets</h2>
<div class="finding"><p class="finding-n">Reason 1 — Performance</p>
<p>XGBoost achieves higher AUC-PR ({xgb_ec["AUCPR"]} vs {lr_ec["AUCPR"]} on e-commerce; {xgb_cc["AUCPR"]} vs {lr_cc["AUCPR"]} on credit card) and higher F1 on both datasets.</p></div>
<div class="finding"><p class="finding-n">Reason 2 — Fewer Missed Frauds</p>
<p>XGBoost produces {int(lr_ec["FN"]) - int(xgb_ec["FN"])} fewer false negatives on e-commerce — fraudulent transactions that would have slipped through undetected with Logistic Regression.</p></div>
<div class="finding"><p class="finding-n">Reason 3 — Cross-Validation Stability</p>
<p>Low CV standard deviation confirms XGBoost generalises reliably. Wide spread would indicate overfitting to SMOTE-augmented samples.</p></div>
<div class="finding"><p class="finding-n">Reason 4 — Non-Linear Fraud Patterns</p>
<p>Fraud involves complex feature interactions (high value AND new account AND high-risk country). XGBoost captures these automatically; Logistic Regression cannot without manual feature engineering.</p></div>
<div class="finding"><p class="finding-n">Reason 5 — SHAP Compatibility</p>
<p>XGBoost integrates directly with SHAP's TreeExplainer — providing fast, exact Shapley values for regulatory-grade explainability in Task 3.</p></div>
<hr>

<h2>10. What's Next — Task 3 (Final)</h2>
<table>
<thead><tr><th>Task</th><th>Detail</th></tr></thead>
<tbody>
<tr><td>SHAP summary plot</td><td>Global feature importance on test set (beeswarm + bar)</td></tr>
<tr><td>SHAP force plots × 3</td><td>One true positive, one false positive, one false negative — with annotation</td></tr>
<tr><td>Built-in vs SHAP</td><td>Validate gain importance against SHAP rank; explain discrepancies</td></tr>
<tr><td>Business recommendations</td><td>≥3 actionable insights connected to specific SHAP values</td></tr>
<tr><td>Final report</td><td>End-to-end blog/PDF covering all three tasks</td></tr>
</tbody>
</table>
<hr>

<p><strong>GitHub:</strong> <a href="https://github.com/maedotamha/Fraud-Cases">github.com/maedotamha/Fraud-Cases</a><br>
Modeling notebook fully executed — all 4 models saved to <code>models/</code>.</p>
<div class="footer">Built with Python 3.11 · XGBoost · scikit-learn · imbalanced-learn · matplotlib · seaborn<br>
Adey Innovations Inc. · Fraud Detection Pipeline · Interim Report 2 · June 2026</div>
</body></html>"""

html_path = OUT / "interim2_medium_post.html"
html_path.write_text(html, encoding="utf-8")
print(f"HTML report saved: {html_path}")
print("\nDone — both reports ready.")
